"""Editable exports from saved evaluations
no financial recalculation or AI calls."""

from io import BytesIO
import math
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Inches, Pt, RGBColor

COLUMNS = [
    ("Manager", "canonical_name"),
    ("Strategy", "strategy"),
    ("Liquidity", "liquidity"),
    ("Eligible", "eligible"),
    ("Rank", "rank"),
    ("Score out of 100", "total_score"),
    ("Annualized return", "annualized_return"),
    ("Annualized volatility", "annualized_volatility"),
    ("Maximum drawdown", "maximum_drawdown"),
    ("Sharpe", "sharpe_ratio"),
    ("Benchmark correlation", "benchmark_correlation"),
    ("Months", "observation_count"),
    ("Period start", "period_start"),
    ("Period end", "period_end"),
]


def safe_cell(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Untrusted names/notes must never become Excel formulas.


def comparison_xlsx(result):
    wb = Workbook()
    ws = wb.active
    ws.title = "Manager comparison"
    ws.append(["Manager comparison for " + result["mandate"]["name"]])
    ws.append(["Saved evaluation " + result["run_id"]])
    ws.append(["As of " + result["created_at"] + " | " + result["market_status"]])
    ws.append(["Static snapshot. Different history windows are shown explicitly. Scores reflect configured policy."])
    headers = [c[0] for c in COLUMNS] + ["Why excluded", "Unverified source concerns", "Decision evidence", "Manager data"]
    ws.append(headers)
    for fund in result["funds"]:
        values = [fund.get(key, fund["metrics"].get(key)) for _, key in COLUMNS]
        values += [
            ", ".join(c["reason"] for c in fund["failed_constraints"]),
            fund["qualitative_notes"],
            fund["decision_evidence"],
            "Synthetic" if fund["synthetic"] else "Uploaded",
        ]
        values[3] = "Yes" if fund["eligible"] else "No"
        for col, value in enumerate(values, 1):
            cell = ws.cell(ws.max_row + (col == 1), col)
            safe_cell(cell, value)
            if 7 <= col <= 9:
                cell.number_format = "0.0%"
            elif col in (6, 10, 11):
                cell.number_format = "0.00"
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:R{ws.max_row}"
    ws.row_dimensions[5].height = 34
    for cell in ws[5]:
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="203739")
    for row in ws.iter_rows(min_row=6):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F4EF")
        ws.row_dimensions[row[0].row].height = 58
    for col in "ABCDEFGHIJKLMNOPQR":
        ws.column_dimensions[col].width = 18
    for col in ("A", "O", "P", "Q"):
        ws.column_dimensions[col].width = 40 if col != "Q" else 65
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["K"].width = 25
    for row in ws.iter_rows(min_row=6):
        longest = max(math.ceil(len(str(cell.value or "")) / ws.column_dimensions[cell.column_letter].width) for cell in row)
        ws.row_dimensions[row[0].row].height = max(42, longest * 15 + 10)
    ws.row_dimensions[4].height = 30
    ws["A4"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A1:R1")
    ws.merge_cells("A2:R2")
    ws.merge_cells("A3:R3")
    ws.merge_cells("A4:R4")
    ws["A1"].font = Font(size=18, bold=True, color="203739")
    ms = wb.create_sheet("Mandate and sources")
    ms.append(["Saved evaluation context", "Value"])
    for key, value in result["mandate"].items():
        display = result["weights"] if key == "weights" else value
        if isinstance(display, dict):
            display = ", ".join(f"{k.replace('_', ' ')} {v:.0%}" for k, v in display.items())
        elif isinstance(display, list):
            display = ", ".join(display) or "None specified"
        elif display is None:
            display = "Not configured"
        ms.append([key.replace("_", " ").capitalize(), display])
        if key in {"maximum_drawdown", "target_volatility", "maximum_volatility"} and isinstance(value, (int, float)):
            ms.cell(ms.max_row, 2).number_format = "0.0%"
    ms.append(["Evaluation ID", result["run_id"]])
    ms.append(["Interpretation", "Synthetic managers are explicitly marked in the app. Market fixture status is recorded below. No allocation recommendation."])
    for key, source in result["sources"].items():
        ms.append([key, f"{source['provider']} | {source['live_or_cached']} | {source['retrieved_at']} | {source['transformation']}"])
    for row in ms:
        for cell in row:
            safe_cell(cell, cell.value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ms.column_dimensions["A"].width = 35
    ms.column_dimensions["B"].width = 100
    for row in range(1, ms.max_row + 1):
        ms.row_dimensions[row].height = 42
    for cell in ms[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="203739")
    ms.freeze_panes = "B2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def memo_docx(result, payload, metadata, evidence, include_evidence=False):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        if style.element.pPr is not None:
            borders = style.element.pPr.find(qn("w:pBdr"))
            if borders is not None:
                style.element.pPr.remove(borders)
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(3)
    doc.styles["Normal"].paragraph_format.line_spacing = 1
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Title"].paragraph_format.space_after = Pt(5)
    doc.styles["Heading 1"].font.size = Pt(11)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(7)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(3)
    doc.add_heading("Investment committee research memo", 0)
    doc.add_paragraph(f"{result['mandate']['name']} | {result['mandate']['benchmark']} | {result['created_at'][:10]}")
    doc.add_paragraph(
        "Purpose: review manager suitability under the saved mandate and identify diligence required before any allocation. This editable document is a snapshot; edits in Word do not update the app or its evidence."
    )
    doc.add_paragraph(f"Evaluation {result['run_id']}\nSource status {result['market_status']}\nGeneration {metadata.get('provider', 'Saved memo')}")
    refs = []
    for item in payload["sections"]:
        doc.add_heading(item["title"].replace("/", "and").replace("-", " "), 1)
        for claim in item["claims"]:
            numbers = []
            for eid in claim["evidence_ids"]:
                if eid not in refs:
                    refs.append(eid)
                numbers.append(str(refs.index(eid) + 1))
            doc.add_paragraph(claim["text"] + (" [" + ", ".join(numbers) + "]" if numbers else ""))
    if include_evidence:
        doc.add_page_break()
    doc.add_heading("Evidence appendix" if include_evidence else "Evidence references", 1)
    if not include_evidence:
        intro = doc.add_paragraph("References below belong to the evaluation above. Full evidence IDs use its ID followed by a colon and the reference suffix. Inspect formulas and original sources in Audit, or download the full evidence export.")
        intro.paragraph_format.keep_with_next = True
    for n, eid in enumerate(refs, 1):
        e = evidence.get(eid)
        if include_evidence:
            doc.add_paragraph(
                f"[{n}] {e.label if e else 'Evidence unavailable'}\n{eid}\n"
                + (f"Source: {e.source}. Period: {e.observation_period}. Method: {e.methodology}" if e else "Inspect the saved evaluation in the app.")
            )
        else:
            suffix = eid.removeprefix(result["run_id"] + ":")
            p = doc.add_paragraph(f"[{n}] {e.label if e else 'Evidence unavailable'} | {suffix}")
            p.paragraph_format.space_after = Pt(1)
            for run in p.runs:
                run.font.size = Pt(8)
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.text.startswith("["):
            paragraph.paragraph_format.keep_together = True
    doc.sections[0].footer.paragraphs[0].text = "Allocator IC Copilot | Internal research prototype"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
