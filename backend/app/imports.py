"""Review imports without side effects; keep original bytes and explicit edits."""

import csv
import hashlib
import json
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from sqlalchemy import select, case
from app.ingestion import ALIASES, ingest_csv
from app.models import Upload, DataQualityIssue

FIELDS = ["fund_id", "fund_name", "manager_name", "strategy", "liquidity", "management_fee", "performance_fee", "period", "monthly_return", "qualitative_notes"]
REQUIRED = ["fund_id", "fund_name", "strategy", "liquidity", "period", "monthly_return"]


def read_table(filename, data, sheet=""):
    if filename.lower().endswith(".csv"):
        content = data.decode("utf-8-sig")
        sheets, selected = [], ""
    elif filename.lower().endswith(".xlsx"):
        try:
            with ZipFile(BytesIO(data)) as z:
                if len(z.infolist()) > 3000 or sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
                    raise ValueError("Workbook expands beyond the supported size")
            wb = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        except (BadZipFile, KeyError, OSError) as exc:
            raise ValueError("Invalid XLSX workbook") from exc
        try:
            sheets = wb.sheetnames
            selected = sheet or sheets[0]
            if selected not in sheets:
                raise ValueError("Worksheet not found")
            ws = wb[selected]
            if ws.max_column > 60 or ws.max_row > 50001:
                raise ValueError("At most 60 columns and 50,000 data rows are supported")
            out = StringIO()
            writer = csv.writer(out)
            for row in ws.iter_rows():
                values = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        # Never execute or silently trust cached workbook formulas.
                        value = "[FORMULA: replace with a verified value]"
                    elif isinstance(value, (date, datetime)):
                        value = value.strftime("%Y-%m-%d")
                    elif isinstance(value, (int, float)) and "%" in cell.number_format:
                        value = f"{value * 100:.12g}%"
                    values.append("" if value is None else str(value))
                writer.writerow(values)
            content = out.getvalue()
        finally:
            wb.close()
    else:
        raise ValueError("Use CSV or XLSX; legacy XLS and macro-enabled files are unsupported")
    table = list(csv.reader(StringIO(content)))
    if not table or not table[0] or len(table[0]) > 60:
        raise ValueError("First row must contain at most 60 column headers")
    headers = table[0]
    if any(not h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("Headers must be nonempty and unique")
    if len(table) > 50001:
        raise ValueError("At most 50,000 data rows are supported")
    suggested = {h: ALIASES.get(re.sub(r"\s+", "_", h.strip().lower()), re.sub(r"\s+", "_", h.strip().lower())) for h in headers}
    suggested = {h: f if f in FIELDS else "" for h, f in suggested.items()}
    return dict(content=content, headers=headers, rows=table[1:], sheets=sheets, sheet=selected, suggested=suggested)


def prepare(table, mapping, corrections):
    mapping = mapping if mapping else table["suggested"]
    if set(mapping) - set(table["headers"]) or any(v and v not in FIELDS for v in mapping.values()):
        raise ValueError("Unknown column or target field in mapping")
    chosen = [(h, mapping.get(h, "")) for h in table["headers"] if mapping.get(h, "")]
    targets = [f for _, f in chosen]
    if len(set(targets)) != len(targets):
        raise ValueError("Each target field must be mapped only once")
    missing = set(REQUIRED) - set(targets)
    if missing:
        raise ValueError("Map required fields: " + ", ".join(sorted(missing)))
    if len(corrections) > 5000:
        raise ValueError("At most 5,000 corrected rows per import")
    for row, edits in corrections.items():
        if not row.isdigit() or not 2 <= int(row) <= len(table["rows"]) + 1:
            raise ValueError("Correction row is outside the worksheet")
        if not isinstance(edits, dict) or set(edits) - set(table["headers"]) or any(not isinstance(v, str) or len(v) > 5000 for v in edits.values()):
            raise ValueError("Corrections must use original columns and text values up to 5,000 characters")
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow(targets)
    for index, original in enumerate(table["rows"], 2):
        if len(original) != len(table["headers"]):
            raise ValueError(f"Row {index} has a different column count; correct the source file")
        values = dict(zip(table["headers"], original))
        values.update(corrections.get(str(index), {}))
        writer.writerow([values[h] for h, _ in chosen])
    return out.getvalue(), mapping


def import_table(db, filename, data, table, mapping, corrections):
    content, mapping = prepare(table, mapping, corrections)
    # Distinct reviewed versions have distinct ingestion identities; original file
    # digest is retained separately, alongside the exact original file bytes.
    context = json.dumps(dict(original_sha256=hashlib.sha256(data).hexdigest(), sheet=table["sheet"], mapping=mapping, corrections=corrections), sort_keys=True)
    summary = ingest_csv(db, filename, content, context)
    upload = db.get(Upload, summary["upload_id"])
    if not summary.get("already_imported"):
        upload.original_file = data
        upload.source_metadata = dict(
            original_sha256=hashlib.sha256(data).hexdigest(), sheet=table["sheet"], mapping=mapping, corrections=corrections, original_table=table["content"]
        )
        if corrections or mapping != table["suggested"]:
            db.add(
                DataQualityIssue(
                    source_file=upload.id,
                    issue_type="reviewed_import",
                    severity="info",
                    description="User confirmed column mapping and corrections before import",
                    inferred_resolution=json.dumps(dict(mapping=mapping, corrections=corrections)),
                    status="resolved",
                )
            )
        if corrections or mapping != table["suggested"]:
            summary = {**summary, "normalization_actions": summary["normalization_actions"] + 1}
            upload.summary = summary
        db.flush()
    return summary


def preview(db, filename, data, table, mapping, corrections):
    # The real ingestion validator runs against current DB state, within a rolled
    # back savepoint. Preview never adds uploads, funds, returns, or issues.
    transaction = db.begin_nested()
    try:
        summary = import_table(db, filename, data, table, mapping, corrections)
        issues = db.scalars(
            select(DataQualityIssue)
            .where(DataQualityIssue.source_file == summary["upload_id"])
            .order_by(case((DataQualityIssue.severity == "error", 0), (DataQualityIssue.severity == "warning", 1), else_=2), DataQualityIssue.source_row)
            .limit(300)
        ).all()
        result = dict(
            summary=summary,
            issues=[dict(source_row=i.source_row, severity=i.severity, description=i.description, issue_type=i.issue_type) for i in issues],
            issues_limited=len(issues) == 300,
        )
    finally:
        transaction.rollback()
    return result
