from copy import deepcopy
from io import BytesIO
import hashlib
import json
import pytest
from openpyxl import Workbook, load_workbook
from docx import Document
from sqlalchemy import select, func
from fastapi.testclient import TestClient
from app.main import app
from app.db import get_db
from app.models import Upload, Fund, MonthlyReturn, Evidence
from app.imports import read_table, prepare, preview, import_table
from app.exports import comparison_xlsx, memo_docx
from app.evidence import briefing_sections, persist_memo
from app.service import seed_demo, evaluate
from app.mandates import MandateInput

CSV = b"Code,Name,Type,Terms,Date,Performance\nREVIEW,Review Macro,macro,monthly,2020-01,2\n"
MAPPING = dict(Code="fund_id", Name="fund_name", Type="strategy", Terms="liquidity", Date="period", Performance="monthly_return")


def test_preview_rolls_back_and_corrections_preserve_original(db):
    table = read_table("review.csv", CSV)
    assert preview(db, "review.csv", CSV, table, MAPPING, {})["summary"]["errors"] == 1
    edits = {"2": {"Performance": "2%"}}
    result = preview(db, "review.csv", CSV, table, MAPPING, edits)
    assert result["summary"]["valid_rows"] == 1
    for model in [Upload, Fund, MonthlyReturn]:
        assert db.scalar(select(func.count()).select_from(model)) == 0
    summary = import_table(db, "review.csv", CSV, table, MAPPING, edits)
    upload = db.get(Upload, summary["upload_id"])
    assert upload.original_file == CSV
    assert upload.source_metadata["original_sha256"] == hashlib.sha256(CSV).hexdigest()
    assert upload.source_metadata["corrections"] == edits
    assert db.scalar(select(MonthlyReturn)).return_value == 0.02
    assert import_table(db, "review.csv", CSV, table, MAPPING, edits)["already_imported"]
    conflicting = import_table(db, "review.csv", CSV, table, MAPPING, {"2": {"Performance": "3%"}})
    assert conflicting["errors"] == 1
    assert db.scalar(select(MonthlyReturn)).return_value == 0.02


def test_xlsx_dates_percentages_sheets_and_formulas():
    from datetime import datetime

    wb = Workbook()
    ws = wb.active
    ws.title = "Managers"
    ws.append(["fund_id", "fund_name", "strategy", "liquidity", "period", "monthly_return"])
    ws.append(["BOOK", "Workbook Macro", "macro", "monthly", datetime(2020, 1, 31), 0.012])
    ws["F2"].number_format = "0.0%"
    ws.append(["BOOK", "Workbook Macro", "macro", "monthly", "2020-02", "=1+1"])
    wb.create_sheet("Other")
    stream = BytesIO()
    wb.save(stream)
    table = read_table("managers.xlsx", stream.getvalue())
    assert table["sheets"] == ["Managers", "Other"]
    assert table["rows"][0][-2:] == ["2020-01-31", "1.2%"]
    assert table["rows"][1][-1].startswith("[FORMULA")
    with pytest.raises(ValueError, match="Worksheet"):
        read_table("managers.xlsx", stream.getvalue(), "Missing")
    with pytest.raises(ValueError, match="CSV or XLSX"):
        read_table("managers.xls", stream.getvalue())
    with pytest.raises(ValueError, match="only once"):
        prepare(table, {**table["suggested"], "strategy": "fund_id"}, {})
    with pytest.raises(ValueError, match="outside"):
        prepare(table, {}, {"999": {"monthly_return": "1%"}})


def test_exports_are_snapshots_and_escape_formula_injection(db):
    seed_demo(db)
    run = evaluate(db, MandateInput())
    saved = deepcopy(run.result)
    export = deepcopy(saved)
    export["funds"][0]["canonical_name"] = '=HYPERLINK("https://example.com")'
    wb = load_workbook(BytesIO(comparison_xlsx(export)))
    ws = wb["Manager comparison"]
    assert ws.max_row == 23
    assert ws["A6"].data_type == "s"
    assert ws["G6"].value == saved["funds"][0]["metrics"]["annualized_return"]
    assert ws["G6"].number_format == "0.0%"
    assert ws.freeze_panes == "C6"
    assert ws.auto_filter.ref == "A5:R23"
    evidence = {e.id: e for e in db.scalars(select(Evidence).where(Evidence.run_id == run.id))}
    payload = briefing_sections(saved)
    document = Document(BytesIO(memo_docx(saved, payload, {"provider": "Rules-based"}, evidence)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert saved["run_id"] in text and "Evidence references" in text
    detailed = Document(BytesIO(memo_docx(saved, payload, {}, evidence, include_evidence=True)))
    assert "Evidence appendix" in "\n".join(p.text for p in detailed.paragraphs)
    for eid in payload["sections"][0]["claims"][0]["evidence_ids"]:
        assert eid.removeprefix(saved["run_id"] + ":") in text
    assert payload["sections"][0]["claims"][0]["text"] in text
    evaluate(db, MandateInput(maximum_drawdown=0.07), run.id)
    assert run.result == saved
    top = next(f for f in saved["funds"] if f["eligible"])
    assert top["suitability"]["strengths"]
    assert top["unresolved_data_quality"] == 0
    assert not top["suitability"]["missing_evidence"]
    assert "Passes" in top["suitability"]["headline"]
    assert next(f for f in saved["funds"] if not f["eligible"])["suitability"]["concerns"]


def test_reviewed_import_api_and_export_run_ownership(db):
    app.dependency_overrides[get_db] = lambda: db
    try:
        with TestClient(app) as client:
            initial = client.post("/api/imports/preview", files={"file": ("review.csv", CSV)}).json()
            assert "validation_error" in initial and initial["row_count"] == 1
            data = {"mapping": json.dumps(MAPPING), "corrections": json.dumps({"2": {"Performance": "2%"}})}
            p = client.post("/api/imports/preview", files={"file": ("review.csv", CSV)}, data=data)
            assert p.status_code == 200 and p.json()["summary"]["valid_rows"] == 1
            assert db.scalar(select(func.count()).select_from(Upload)) == 0
            seed_demo(db)
            committed = client.post("/api/imports/commit", files={"file": ("review.csv", CSV)}, data=data)
            assert committed.status_code == 200
            result = committed.json()
            row = client.get(f"/api/uploads/{result['summary']['upload_id']}/rows/2").json()
            assert client.get(f"/api/uploads/{result['summary']['upload_id']}/original").content == CSV
            assert row["values"][-1] == "2" and row["corrections"]["Performance"] == "2%"
            rid = result["result"]["run_id"]
            for ext in ["comparison.xlsx", "memo.docx"]:
                r = client.get(f"/api/runs/{rid}/{ext}")
                assert r.status_code == 200 and r.content.startswith(b"PK")
            detailed = client.get(f"/api/runs/{rid}/memo.docx?include_evidence=true")
            assert detailed.status_code == 200
            assert "committee-evidence" in detailed.headers["content-disposition"]
            assert "Evidence appendix" in "\n".join(p.text for p in Document(BytesIO(detailed.content)).paragraphs)
            other = evaluate(db, MandateInput())
            memo = persist_memo(db, other, briefing_sections(other.result), "briefing", {})
            assert client.get(f"/api/runs/{rid}/memo.docx?memo_id={memo.id}").status_code == 404
            bad = client.post("/api/imports/commit", files={"file": ("review.csv", CSV)}, data={"mapping": "[]"})
            assert bad.status_code == 422
    finally:
        app.dependency_overrides.clear()
